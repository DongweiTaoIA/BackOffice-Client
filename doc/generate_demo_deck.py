"""Generate the BackOffice Admin Portal demo deck (visual + business-focused).

Run with:  python doc/generate_demo_deck.py
Outputs:   doc/BackOffice-Admin-Portal-Demo.pptx

Design intent: pictures, shapes and emoji over paragraphs. Less text, more
visual story-telling. Aligned with doc/demo-script.md.
"""

import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# ---------------------------------------------------------------------------
# Brand palette
# ---------------------------------------------------------------------------
PRIMARY = RGBColor(0x4F, 0x46, 0xE5)        # indigo-600
PRIMARY_DARK = RGBColor(0x31, 0x2E, 0x81)   # indigo-900
PRIMARY_LIGHT = RGBColor(0xC7, 0xD2, 0xFE)  # indigo-200
ACCENT_PINK = RGBColor(0xEC, 0x48, 0x99)
ACCENT_TEAL = RGBColor(0x14, 0xB8, 0xA6)
ACCENT_ORANGE = RGBColor(0xF9, 0x73, 0x16)
ACCENT_GREEN = RGBColor(0x22, 0xC5, 0x5E)
ACCENT_YELLOW = RGBColor(0xFA, 0xCC, 0x15)
ACCENT_RED = RGBColor(0xEF, 0x44, 0x44)

TEXT_DARK = RGBColor(0x11, 0x18, 0x27)
TEXT_MUTED = RGBColor(0x6B, 0x72, 0x80)
BG_LIGHT = RGBColor(0xF9, 0xFA, 0xFB)
BG_GRAY = RGBColor(0xE5, 0xE7, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def add_bg(slide, color=BG_LIGHT):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    return bg


def add_top_bar(slide, color=PRIMARY):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.18))
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = color


def add_title(slide, text, color=PRIMARY_DARK, size=34, top=Inches(0.45),
              left=Inches(0.6), width=Inches(12.1), height=Inches(0.9),
              align=PP_ALIGN.LEFT, bold=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_subtitle(slide, text, color=TEXT_MUTED, size=16, top=Inches(1.2),
                 left=Inches(0.6), width=Inches(12.1), align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.color.rgb = color
    return box


def add_text(slide, text, left, top, width, height, *, size=14, color=TEXT_DARK,
             bold=False, italic=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(text, str):
        lines = text.split("\n")
    else:
        lines = list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.italic = italic
        p.font.color.rgb = color
    return box


def add_rounded(slide, left, top, width, height, *, fill=WHITE,
                line_color=None, line_width=None, corner=0.08):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    s.adjustments[0] = corner
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line_color is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_color
        if line_width is not None:
            s.line.width = line_width
    return s


def add_pill(slide, left, top, width, height, label, *, fill=PRIMARY,
             font_color=WHITE, size=12, bold=True):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  left, top, width, height)
    pill.adjustments[0] = 0.5
    pill.line.fill.background()
    pill.fill.solid()
    pill.fill.fore_color.rgb = fill
    tf = pill.text_frame
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = label
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = font_color
    return pill


def add_circle(slide, cx, cy, radius, *, fill=PRIMARY, line_color=None):
    left = cx - radius
    top = cy - radius
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, radius * 2, radius * 2)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line_color is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line_color
    return s


def add_emoji(slide, emoji, left, top, width, height, *, size=48,
              color=TEXT_DARK, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = emoji
    p.font.size = Pt(size)
    p.font.color.rgb = color
    return box


def add_footer(slide, page_no, total):
    add_text(slide, "BackOffice Admin Portal  \u00b7  Demo",
             Inches(0.4), Inches(7.15), Inches(7), Inches(0.3),
             size=10, color=TEXT_MUTED)
    add_text(slide, f"{page_no} / {total}",
             Inches(12.0), Inches(7.15), Inches(1.0), Inches(0.3),
             size=10, color=TEXT_MUTED, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY_DARK

    # Diagonal accent block.
    accent = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE,
                                    Inches(8.0), Inches(0), Inches(5.33), SLIDE_H)
    accent.line.fill.background()
    accent.fill.solid()
    accent.fill.fore_color.rgb = PRIMARY
    accent.rotation = 180

    # Floating circles for visual interest.
    add_circle(slide, Inches(10.5), Inches(1.4), Inches(0.45), fill=ACCENT_PINK)
    add_circle(slide, Inches(11.6), Inches(2.0), Inches(0.25), fill=ACCENT_TEAL)
    add_circle(slide, Inches(12.4), Inches(6.0), Inches(0.6), fill=ACCENT_YELLOW)
    add_circle(slide, Inches(9.6), Inches(6.5), Inches(0.3), fill=PRIMARY_LIGHT)

    add_pill(slide, Inches(0.9), Inches(1.4), Inches(2.4), Inches(0.45),
             "\U0001F4A1  INNOVATION PITCH", fill=ACCENT_PINK, size=12)

    add_text(slide, "BackOffice", Inches(0.9), Inches(2.0), Inches(11), Inches(1.4),
             size=68, bold=True, color=WHITE)
    add_text(slide, "Admin Portal.", Inches(0.9), Inches(3.0), Inches(11), Inches(1.4),
             size=68, bold=True, color=ACCENT_PINK)

    add_text(slide,
             "One admin-only portal.  Self-service.  AI inside.",
             Inches(0.9), Inches(4.4), Inches(11), Inches(0.6),
             size=22, color=PRIMARY_LIGHT)

    add_emoji(slide, "\U0001F680  \U0001F916  \U0001F4CA  \U0001F6E1",
              Inches(0.9), Inches(5.4), Inches(7), Inches(1),
              size=44, color=WHITE, align=PP_ALIGN.LEFT)


def slide_problem(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide, ACCENT_RED)
    add_title(slide, "Today's reality \U0001F629")
    add_subtitle(slide, "Admins juggle aging systems and constant IT tickets",
                 color=TEXT_MUTED)

    systems = [
        ("Unifi",         Inches(0.8),  Inches(2.0),  ACCENT_ORANGE, -6),
        ("Unification",   Inches(3.1),  Inches(2.5),  ACCENT_TEAL,    4),
        ("TPA",           Inches(5.6),  Inches(2.0),  PRIMARY,       -3),
        ("Batch Jobs",    Inches(7.8),  Inches(2.6),  ACCENT_PINK,    7),
        ("Payments",      Inches(10.4), Inches(2.1),  ACCENT_YELLOW, -5),
    ]
    for name, left, top, color, rot in systems:
        card = add_rounded(slide, left, top, Inches(2.1), Inches(1.2),
                           fill=color, corner=0.18)
        card.rotation = rot
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = name
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = WHITE

    pains = [
        ("\U0001F512", "Shared with external users",
         "Admin & external in the SAME UI = security risk"),
        ("\U0001F39F",  "IT-ticket bottleneck",
         "Program setup, operator setup, data fix \u2192 wait days"),
        ("\U0001F300",  "5+ systems to juggle",
         "Each with its own login, UI and quirks"),
    ]
    card_w = Inches(4.0)
    card_h = Inches(2.0)
    gap = Inches(0.2)
    start_left = Inches(0.5)
    top = Inches(4.6)
    for i, (emoji, head, sub) in enumerate(pains):
        left = start_left + (card_w + gap) * i
        add_rounded(slide, left, top, card_w, card_h,
                    fill=WHITE, line_color=BG_GRAY, line_width=Pt(0.75),
                    corner=0.1)
        add_emoji(slide, emoji, left + Inches(0.2), top + Inches(0.25),
                  Inches(0.9), Inches(0.9), size=36)
        add_text(slide, head, left + Inches(1.15), top + Inches(0.3),
                 card_w - Inches(1.3), Inches(0.4),
                 size=15, bold=True, color=TEXT_DARK)
        add_text(slide, sub, left + Inches(1.15), top + Inches(0.85),
                 card_w - Inches(1.3), Inches(1.0),
                 size=11, color=TEXT_MUTED)

    add_footer(slide, page, total)


def slide_vision(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Tomorrow: one front door \U0001F6AA")
    add_subtitle(slide, "One admin-only portal \u00b7 every system, one screen")

    # Center hub.
    center_x = Inches(6.667)
    center_y = Inches(4.2)
    hub_r = Inches(1.4)
    hub = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        center_x - hub_r, center_y - hub_r, hub_r * 2, hub_r * 2,
    )
    hub.fill.solid()
    hub.fill.fore_color.rgb = PRIMARY
    hub.line.color.rgb = PRIMARY_DARK
    hub.line.width = Pt(2)
    tf = hub.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "BackOffice\nAdmin Portal"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE

    spokes = [
        ("Unifi",        ACCENT_ORANGE),
        ("Unification",  ACCENT_TEAL),
        ("TPA",          PRIMARY_LIGHT),
        ("Batch Jobs",   ACCENT_PINK),
        ("Payments",     ACCENT_YELLOW),
    ]
    orbit_emu = int(Inches(2.8))
    spoke_w = Inches(1.8)
    spoke_h = Inches(0.7)
    for i, (name, color) in enumerate(spokes):
        angle = math.radians(-90 + i * (360 / len(spokes)))
        cx = int(center_x) + int(orbit_emu * math.cos(angle))
        cy = int(center_y) + int(orbit_emu * math.sin(angle))
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(int(cx - spoke_w / 2)),
            Emu(int(cy - spoke_h / 2)),
            spoke_w, spoke_h,
        )
        pill.adjustments[0] = 0.5
        pill.fill.solid()
        pill.fill.fore_color.rgb = color
        pill.line.fill.background()
        ptf = pill.text_frame
        ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = ptf.paragraphs[0]
        pp.alignment = PP_ALIGN.CENTER
        pp.text = name
        pp.font.size = Pt(13)
        pp.font.bold = True
        pp.font.color.rgb = WHITE

        # Connector line from hub to spoke.
        line = slide.shapes.add_connector(1, center_x, center_y, Emu(cx), Emu(cy))
        line.line.color.rgb = PRIMARY_LIGHT
        line.line.width = Pt(2)

    badges = [
        ("\U0001F6E1",   "Secure",       ACCENT_GREEN),
        ("\u26A1",        "Self-service", ACCENT_PINK),
        ("\U0001F310",    "Unified",      PRIMARY),
    ]
    bw = Inches(2.6)
    bh = Inches(0.65)
    bgap = Inches(0.4)
    btotal = bw * 3 + bgap * 2
    bstart = (SLIDE_W - btotal) / 2
    btop = Inches(6.35)
    for i, (emoji, label, color) in enumerate(badges):
        left = bstart + (bw + bgap) * i
        add_pill(slide, left, btop, bw, bh, f"{emoji}   {label}",
                 fill=color, size=16)

    add_footer(slide, page, total)


def slide_why_now(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Why now? \U0001F525")
    add_subtitle(slide, "Three forces line up at the same moment")

    pillars = [
        ("\u26A1", "EFFICIENCY",
         "Collapse multi-system workflows.\nKill the IT-ticket queue.",
         ACCENT_PINK),
        ("\U0001F6E1", "SECURITY",
         "Admin tooling, finally separated\nfrom external-facing systems.",
         PRIMARY),
        ("\U0001F916", "AI",
         "LLMs are good enough now to act\nas a real co-pilot for admins.",
         ACCENT_TEAL),
    ]
    card_w = Inches(3.9)
    card_h = Inches(4.6)
    gap = Inches(0.3)
    total_w = card_w * 3 + gap * 2
    start_left = (SLIDE_W - total_w) / 2
    top = Inches(2.0)

    for i, (emoji, head, body, color) in enumerate(pillars):
        left = start_left + (card_w + gap) * i
        add_rounded(slide, left, top, card_w, card_h, fill=WHITE,
                    line_color=BG_GRAY, line_width=Pt(0.75), corner=0.06)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top,
                                      card_w, Inches(1.6))
        band.line.fill.background()
        band.fill.solid()
        band.fill.fore_color.rgb = color
        add_emoji(slide, emoji, left, top + Inches(0.25),
                  card_w, Inches(1.1), size=64, color=WHITE)
        add_text(slide, head, left, top + Inches(1.8),
                 card_w, Inches(0.6),
                 size=22, bold=True, color=TEXT_DARK,
                 align=PP_ALIGN.CENTER)
        add_text(slide, body, left + Inches(0.3), top + Inches(2.6),
                 card_w - Inches(0.6), Inches(1.8),
                 size=14, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    add_footer(slide, page, total)


def slide_what(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "What is it? \U0001F9D0")
    add_subtitle(slide, "A modern, admin-only web console")

    quads = [
        ("\U0001F310", "Web-based",
         "No installs. Any laptop. Always up to date.", PRIMARY),
        ("\U0001F511", "Microsoft sign-in",
         "Single sign-on with Entra ID \u2014 only admins.", ACCENT_TEAL),
        ("\U0001F4CA", "Live dashboard",
         "Real-time pulse across every connected system.", ACCENT_PINK),
        ("\U0001F3AF", "Search \u2192 drill \u2192 act",
         "Same pattern in every module. Learn one, know all.", ACCENT_ORANGE),
    ]
    card_w = Inches(5.7)
    card_h = Inches(2.1)
    gap_x = Inches(0.3)
    gap_y = Inches(0.3)
    total_w = card_w * 2 + gap_x
    start_left = (SLIDE_W - total_w) / 2
    start_top = Inches(2.0)

    for i, (emoji, head, body, color) in enumerate(quads):
        row = i // 2
        col = i % 2
        left = start_left + (card_w + gap_x) * col
        top = start_top + (card_h + gap_y) * row
        add_rounded(slide, left, top, card_w, card_h,
                    fill=WHITE, line_color=BG_GRAY, line_width=Pt(0.75),
                    corner=0.08)
        add_circle(slide, left + Inches(0.9), top + Inches(1.05),
                   Inches(0.55), fill=color)
        add_emoji(slide, emoji,
                  left + Inches(0.35), top + Inches(0.5),
                  Inches(1.1), Inches(1.1), size=32, color=WHITE)
        add_text(slide, head,
                 left + Inches(1.75), top + Inches(0.45),
                 card_w - Inches(1.9), Inches(0.5),
                 size=20, bold=True, color=TEXT_DARK)
        add_text(slide, body,
                 left + Inches(1.75), top + Inches(1.05),
                 card_w - Inches(1.9), Inches(1.0),
                 size=13, color=TEXT_MUTED)

    add_footer(slide, page, total)


def slide_better_way(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "A better way to work \U0001F680")
    add_subtitle(slide, "Routine work moves from IT tickets into the admins' own hands")

    left_x = Inches(0.6)
    left_w = Inches(5.4)
    col_top = Inches(2.0)
    col_h = Inches(4.6)
    add_rounded(slide, left_x, col_top, left_w, col_h,
                fill=RGBColor(0xF3, 0xF4, 0xF6),
                line_color=BG_GRAY, line_width=Pt(0.75), corner=0.05)
    add_text(slide, "TODAY  \U0001F629",
             left_x + Inches(0.3), col_top + Inches(0.3),
             left_w - Inches(0.6), Inches(0.5),
             size=18, bold=True, color=TEXT_MUTED)

    today_items = [
        ("\U0001F422", "Slow, clunky, aging UI"),
        ("\U0001F512", "Shared with external users"),
        ("\U0001F39F",  "Every setup task \u2192 IT ticket"),
        ("\u23F3",       "Days of waiting for a data fix"),
        ("\U0001F300",  "Jump across 5+ systems"),
    ]
    row_top = col_top + Inches(1.0)
    for emoji, label in today_items:
        add_emoji(slide, emoji, left_x + Inches(0.3), row_top,
                  Inches(0.5), Inches(0.5), size=22)
        add_text(slide, label,
                 left_x + Inches(1.0), row_top + Inches(0.05),
                 left_w - Inches(1.2), Inches(0.5),
                 size=14, color=TEXT_DARK)
        row_top += Inches(0.7)

    right_x = Inches(7.3)
    right_w = Inches(5.4)
    add_rounded(slide, right_x, col_top, right_w, col_h,
                fill=WHITE,
                line_color=PRIMARY, line_width=Pt(1.5), corner=0.05)
    add_text(slide, "WITH THE PORTAL  \U0001F680",
             right_x + Inches(0.3), col_top + Inches(0.3),
             right_w - Inches(0.6), Inches(0.5),
             size=18, bold=True, color=PRIMARY_DARK)

    after_items = [
        ("\u26A1",       "Clean, modern, admin-only UI"),
        ("\U0001F6E1",  "Separate from any external app"),
        ("\u2705",       "Self-service for routine setup"),
        ("\u23F1",       "Minutes, not days"),
        ("\U0001F310",  "One sidebar \u2014 every system"),
    ]
    row_top = col_top + Inches(1.0)
    for emoji, label in after_items:
        add_emoji(slide, emoji, right_x + Inches(0.3), row_top,
                  Inches(0.5), Inches(0.5), size=22)
        add_text(slide, label,
                 right_x + Inches(1.0), row_top + Inches(0.05),
                 right_w - Inches(1.2), Inches(0.5),
                 size=14, bold=True, color=TEXT_DARK)
        row_top += Inches(0.7)

    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(6.05), Inches(4.0), Inches(1.2), Inches(0.7),
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = ACCENT_PINK
    arrow.line.fill.background()

    add_footer(slide, page, total)


def slide_before_after(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Before vs After")
    add_subtitle(slide, "Same job, two very different days")

    before_x = Inches(0.6)
    before_w = Inches(5.4)
    before_top = Inches(2.0)
    before_h = Inches(4.6)
    add_rounded(slide, before_x, before_top, before_w, before_h,
                fill=RGBColor(0xF3, 0xF4, 0xF6),
                line_color=BG_GRAY, line_width=Pt(0.75), corner=0.05)
    add_text(slide, "BEFORE",
             before_x, before_top + Inches(0.25),
             before_w, Inches(0.5),
             size=20, bold=True, color=TEXT_MUTED,
             align=PP_ALIGN.CENTER)

    systems = ["Unifi", "Unification", "TPA", "Batch Jobs", "Payments"]
    sx = before_x + Inches(0.6)
    sy = before_top + Inches(1.0)
    sw = Inches(1.9)
    sh = Inches(0.55)
    colors = [ACCENT_ORANGE, ACCENT_TEAL, PRIMARY, ACCENT_PINK, ACCENT_YELLOW]
    for i, name in enumerate(systems):
        col = i % 2
        row = i // 2
        left = sx + (sw + Inches(0.2)) * col
        top = sy + Inches(0.65) * row
        add_pill(slide, left, top, sw, sh, name, fill=colors[i], size=12)

    add_pill(slide, before_x + Inches(0.6), before_top + Inches(3.6),
             before_w - Inches(1.2), Inches(0.65),
             "\U0001F39F  IT ticket  \u2192  wait days",
             fill=ACCENT_RED, size=14)

    after_x = Inches(7.3)
    after_w = Inches(5.4)
    add_rounded(slide, after_x, before_top, after_w, before_h,
                fill=WHITE, line_color=PRIMARY, line_width=Pt(1.5), corner=0.05)
    add_text(slide, "AFTER",
             after_x, before_top + Inches(0.25),
             after_w, Inches(0.5),
             size=20, bold=True, color=PRIMARY_DARK,
             align=PP_ALIGN.CENTER)

    add_pill(slide, after_x + Inches(0.6), before_top + Inches(1.3),
             after_w - Inches(1.2), Inches(1.2),
             "BackOffice Admin Portal", fill=PRIMARY, size=22)

    chips = [
        ("\u2705", "Self-service"),
        ("\u26A1", "Minutes"),
        ("\U0001F6E1", "Secure"),
    ]
    cw = (after_w - Inches(1.6)) / 3
    cy = before_top + Inches(3.0)
    for i, (emoji, label) in enumerate(chips):
        left = after_x + Inches(0.6) + (cw + Inches(0.2)) * i
        add_pill(slide, left, cy, cw, Inches(0.55),
                 f"{emoji}  {label}", fill=ACCENT_GREEN, size=13)

    add_pill(slide, after_x + Inches(0.6), before_top + Inches(3.85),
             after_w - Inches(1.2), Inches(0.7),
             "\U0001F60A  Happy admins", fill=ACCENT_PINK, size=16)

    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(6.05), Inches(4.0), Inches(1.2), Inches(0.7),
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = ACCENT_PINK
    arrow.line.fill.background()

    add_footer(slide, page, total)


def slide_dashboard(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Real-time pulse \U0001F4CA")
    add_subtitle(slide, "Six live counters from Unifi \u00b7 Unification preview \u00b7 more coming")

    metrics = [
        ("\U0001F3E2", "Dealers",   "3,581",   PRIMARY),
        ("\U0001F4E6", "Products",  "2.7 M",   ACCENT_GREEN),
        ("\U0001F4CB", "Programs",  "185",     ACCENT_TEAL),
        ("\U0001F4C4", "Contracts", "812,969", ACCENT_ORANGE),
        ("\U0001F465", "Customers", "873,325", ACCENT_PINK),
        ("\U0001F697", "Vehicles",  "812,884", PRIMARY_LIGHT),
    ]
    card_w = Inches(3.95)
    card_h = Inches(1.8)
    gap = Inches(0.2)
    total_w = card_w * 3 + gap * 2
    start_left = (SLIDE_W - total_w) / 2
    start_top = Inches(2.05)

    for i, (emoji, label, value, color) in enumerate(metrics):
        row = i // 3
        col = i % 3
        left = start_left + (card_w + gap) * col
        top = start_top + (card_h + gap) * row
        add_rounded(slide, left, top, card_w, card_h, fill=WHITE,
                    line_color=BG_GRAY, line_width=Pt(0.75), corner=0.06)
        add_rounded(slide, left + Inches(0.3), top + Inches(0.35),
                    Inches(1.1), Inches(1.1),
                    fill=color, corner=0.2)
        add_emoji(slide, emoji,
                  left + Inches(0.3), top + Inches(0.35),
                  Inches(1.1), Inches(1.1), size=36, color=WHITE)
        add_text(slide, label,
                 left + Inches(1.6), top + Inches(0.4),
                 card_w - Inches(1.8), Inches(0.45),
                 size=14, color=TEXT_MUTED)
        add_text(slide, value,
                 left + Inches(1.6), top + Inches(0.85),
                 card_w - Inches(1.8), Inches(0.7),
                 size=28, bold=True, color=TEXT_DARK)

    add_pill(slide, Inches(0.6), Inches(6.4), Inches(4.6), Inches(0.5),
             "\U0001F534  Live \u00b7 auto-refresh \u00b7 activity pulses",
             fill=ACCENT_RED, size=12)
    add_pill(slide, Inches(5.4), Inches(6.4), Inches(3.6), Inches(0.5),
             "\U0001F538  Unification preview lane",
             fill=ACCENT_TEAL, size=12)
    add_pill(slide, Inches(9.2), Inches(6.4), Inches(3.6), Inches(0.5),
             "\u2795  TPA \u00b7 Batches \u00b7 Payments next",
             fill=PRIMARY, size=12)

    add_footer(slide, page, total)


def slide_ecosystem(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "One sidebar.  Everything. \U0001F310")
    add_subtitle(slide, "Live now and on the runway")

    col_w = Inches(6.1)
    col_h = Inches(4.7)
    col_top = Inches(2.0)
    gap = Inches(0.3)
    left_x = (SLIDE_W - col_w * 2 - gap) / 2
    right_x = left_x + col_w + gap

    add_rounded(slide, left_x, col_top, col_w, col_h, fill=WHITE,
                line_color=ACCENT_GREEN, line_width=Pt(1.5), corner=0.05)
    add_pill(slide, left_x + Inches(0.3), col_top + Inches(0.3),
             Inches(2.4), Inches(0.5),
             "\u2705  LIVE NOW", fill=ACCENT_GREEN, size=14)
    live_items = [
        ("\U0001F3E2", "Dealers"),
        ("\U0001F4C4", "Contracts"),
        ("\U0001F6E0", "Claims"),
        ("\u2716",      "Cancellations"),
        ("\U0001F4CA", "Live Dashboard"),
        ("\U0001F916", "AI Chat"),
    ]
    iy = col_top + Inches(1.1)
    for emoji, label in live_items:
        add_emoji(slide, emoji, left_x + Inches(0.4), iy,
                  Inches(0.5), Inches(0.45), size=22)
        add_text(slide, label, left_x + Inches(1.05), iy + Inches(0.05),
                 col_w - Inches(1.3), Inches(0.45),
                 size=15, bold=True, color=TEXT_DARK)
        iy += Inches(0.55)

    add_rounded(slide, right_x, col_top, col_w, col_h, fill=WHITE,
                line_color=PRIMARY, line_width=Pt(1.5), corner=0.05)
    add_pill(slide, right_x + Inches(0.3), col_top + Inches(0.3),
             Inches(2.6), Inches(0.5),
             "\U0001F6A7  COMING NEXT", fill=PRIMARY, size=14)
    next_items = [
        ("\U0001F9FE", "Statements"),
        ("\U0001F3DB", "TPA  \u00b7  Batch Jobs  \u00b7  Payments"),
        ("\U0001F6E0", "Program Setup  \u00b7  Operator Setup"),
        ("\U0001F511", "Roles & Permissions"),
        ("\U0001F4E4", "Data Tools (upload \u00b7 fix \u00b7 export)"),
        ("\U0001F4C8", "Reports & Audit  \u00b7  Configuration"),
    ]
    iy = col_top + Inches(1.1)
    for emoji, label in next_items:
        add_emoji(slide, emoji, right_x + Inches(0.4), iy,
                  Inches(0.5), Inches(0.45), size=22)
        add_text(slide, label, right_x + Inches(1.05), iy + Inches(0.05),
                 col_w - Inches(1.3), Inches(0.45),
                 size=14, color=TEXT_DARK)
        iy += Inches(0.55)

    add_footer(slide, page, total)


def slide_ai(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "AI co-pilot \U0001F916")
    add_subtitle(slide, "Ask in plain language \u00b7 the assistant answers, navigates and acts")

    chat_x = Inches(0.7)
    chat_w = Inches(8.4)
    chat_top = Inches(2.0)

    def user_bubble(top, text):
        bub = add_rounded(slide, chat_x + Inches(2.0), top,
                          chat_w - Inches(2.0), Inches(0.7),
                          fill=PRIMARY, corner=0.25)
        tf = bub.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(15)
        p.font.color.rgb = WHITE
        add_emoji(slide, "\U0001F9D1",
                  chat_x + Inches(1.3), top,
                  Inches(0.6), Inches(0.7), size=28)

    def ai_bubble(top, text, h=Inches(0.7)):
        bub = add_rounded(slide, chat_x, top, chat_w - Inches(2.0), h,
                          fill=WHITE, line_color=BG_GRAY,
                          line_width=Pt(0.75), corner=0.18)
        tf = bub.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        add_emoji(slide, "\U0001F916",
                  chat_x + chat_w - Inches(1.9), top,
                  Inches(0.6), Inches(0.7), size=28)

    user_bubble(chat_top, "Find dealers in Quebec")
    ai_bubble(chat_top + Inches(0.9),
              "Found 42 dealers. Opening the list \u2192")

    user_bubble(chat_top + Inches(1.8), "Can contract C-88421 be cancelled?")
    ai_bubble(chat_top + Inches(2.7),
              "\u2705  Yes. Estimated refund: $1,247.85. Start cancellation?")

    user_bubble(chat_top + Inches(3.6), "Set up a new program from template")
    ai_bubble(chat_top + Inches(4.5),
              "Opening Program Setup with the template pre-filled.")

    rx = Inches(9.4)
    rw = Inches(3.4)
    ry = Inches(2.0)
    add_rounded(slide, rx, ry, rw, Inches(4.6),
                fill=PRIMARY_DARK, corner=0.06)
    add_text(slide, "Why it matters",
             rx + Inches(0.3), ry + Inches(0.25),
             rw - Inches(0.6), Inches(0.5),
             size=16, bold=True, color=WHITE)
    bullets = [
        ("\U0001F50E", "Find anything, fast"),
        ("\U0001F9E0", "Answers, not menus"),
        ("\U0001FA84", "Kicks off workflows"),
        ("\u23F1",      "No more IT-ticket wait"),
    ]
    by = ry + Inches(1.0)
    for emoji, line in bullets:
        add_emoji(slide, emoji, rx + Inches(0.25), by,
                  Inches(0.5), Inches(0.5), size=22)
        add_text(slide, line, rx + Inches(0.85), by + Inches(0.05),
                 rw - Inches(1.0), Inches(0.5),
                 size=13, color=WHITE)
        by += Inches(0.8)

    add_footer(slide, page, total)


def slide_feedback(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Built with admins, not for them \U0001F4AC")
    add_subtitle(slide, "Feedback baked into the product")

    items = [
        ("\U0001F41B", "Report a Bug",
         "Rich-text \u00b7 screenshots \u00b7 severity\nstraight into our backlog.",
         ACCENT_PINK),
        ("\U0001F4A1", "Request a Feature",
         "Pitch new capabilities with\ncontext and priority.",
         ACCENT_YELLOW),
        ("\U0001F198", "Ask for Help",
         "Short-form questions answered\nin-thread by the team.",
         ACCENT_TEAL),
    ]
    card_w = Inches(3.9)
    card_h = Inches(3.6)
    gap = Inches(0.3)
    total_w = card_w * 3 + gap * 2
    start_left = (SLIDE_W - total_w) / 2
    top = Inches(2.0)

    for i, (emoji, head, body, color) in enumerate(items):
        left = start_left + (card_w + gap) * i
        add_rounded(slide, left, top, card_w, card_h, fill=WHITE,
                    line_color=BG_GRAY, line_width=Pt(0.75), corner=0.08)
        add_circle(slide, left + card_w / 2, top + Inches(0.95),
                   Inches(0.7), fill=color)
        add_emoji(slide, emoji, left, top + Inches(0.4),
                  card_w, Inches(1.1), size=44, color=WHITE)
        add_text(slide, head, left, top + Inches(1.95),
                 card_w, Inches(0.5),
                 size=20, bold=True, color=TEXT_DARK, align=PP_ALIGN.CENTER)
        add_text(slide, body, left + Inches(0.3), top + Inches(2.55),
                 card_w - Inches(0.6), Inches(1.0),
                 size=13, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    add_pill(slide, Inches(3.5), Inches(6.0), Inches(6.3), Inches(0.6),
             "\U0001F501   We listen \u2192 we ship \u2192 we improve",
             fill=PRIMARY_DARK, size=15)

    add_footer(slide, page, total)


def slide_impact(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Wins for the admin team \U0001F3C6")
    add_subtitle(slide, "Direction of travel \u2014 what we expect to see")

    wins = [
        ("\U0001F4C9", "Security risk",      "DOWN",  ACCENT_GREEN),
        ("\U0001F4C9", "IT tickets",         "DOWN",  ACCENT_GREEN),
        ("\U0001F4C9", "Onboarding time",    "DOWN",  ACCENT_GREEN),
        ("\U0001F4C9", "Operational errors", "DOWN",  ACCENT_GREEN),
        ("\U0001F4C8", "Admin happiness",    "UP",    ACCENT_PINK),
    ]
    card_w = Inches(2.4)
    card_h = Inches(2.9)
    gap = Inches(0.2)
    total_w = card_w * 5 + gap * 4
    start_left = (SLIDE_W - total_w) / 2
    top = Inches(2.2)

    for i, (emoji, label, dir_, color) in enumerate(wins):
        left = start_left + (card_w + gap) * i
        add_rounded(slide, left, top, card_w, card_h, fill=WHITE,
                    line_color=BG_GRAY, line_width=Pt(0.75), corner=0.08)
        add_emoji(slide, emoji, left, top + Inches(0.2),
                  card_w, Inches(1.2), size=52, color=color)
        add_pill(slide, left + Inches(0.5), top + Inches(1.45),
                 card_w - Inches(1.0), Inches(0.45),
                 dir_, fill=color, size=12)
        add_text(slide, label, left + Inches(0.15), top + Inches(2.05),
                 card_w - Inches(0.3), Inches(0.7),
                 size=13, bold=True, color=TEXT_DARK, align=PP_ALIGN.CENTER)

    add_pill(slide, Inches(3.0), Inches(5.7), Inches(7.3), Inches(0.65),
             "Less waiting.  Less back-and-forth.  More work getting done.",
             fill=PRIMARY, size=15)

    add_footer(slide, page, total)


def slide_roadmap(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Roadmap \U0001F6E3")
    add_subtitle(slide, "Now \u00b7 Next \u00b7 Later")

    lanes = [
        ("NOW",   "\u2705",       ACCENT_GREEN, [
            "Live Dashboard",
            "Dealer Servicing",
            "AI Chat",
            "In-product feedback",
        ]),
        ("NEXT",  "\U0001F6A7",   PRIMARY, [
            "Program Setup",
            "Operator Setup",
            "Roles & Permissions",
            "Data Tools",
            "TPA \u00b7 Batch \u00b7 Payments",
        ]),
        ("LATER", "\U0001F52E",   ACCENT_PINK, [
            "Unification cut-over",
            "Deeper cross-system AI",
            "In-app documentation",
        ]),
    ]
    lane_w = Inches(4.0)
    lane_h = Inches(4.6)
    gap = Inches(0.3)
    total_w = lane_w * 3 + gap * 2
    start_left = (SLIDE_W - total_w) / 2
    top = Inches(2.0)

    for i, (title, emoji, color, items) in enumerate(lanes):
        left = start_left + (lane_w + gap) * i
        add_rounded(slide, left, top, lane_w, lane_h, fill=WHITE,
                    line_color=BG_GRAY, line_width=Pt(0.75), corner=0.06)
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                        left, top, lane_w, Inches(1.0))
        header.line.fill.background()
        header.fill.solid()
        header.fill.fore_color.rgb = color
        add_text(slide, f"{emoji}   {title}",
                 left, top + Inches(0.2),
                 lane_w, Inches(0.7),
                 size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        iy = top + Inches(1.3)
        for item in items:
            add_text(slide, f"\u00b7  {item}",
                     left + Inches(0.45), iy,
                     lane_w - Inches(0.6), Inches(0.45),
                     size=14, color=TEXT_DARK)
            iy += Inches(0.55)

    for i in range(2):
        ax = start_left + lane_w + (lane_w + gap) * i - Inches(0.05)
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW, ax, top + Inches(2.0),
            gap + Inches(0.1), Inches(0.45),
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = PRIMARY_LIGHT
        arrow.line.fill.background()

    add_footer(slide, page, total)


def slide_demo(prs, page, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide)
    add_title(slide, "Demo time \U0001F3AC")
    add_subtitle(slide, "A guided 3-minute walkthrough")

    steps = [
        ("\U0001F511", "Sign in"),
        ("\U0001F4CA", "Dashboard"),
        ("\U0001F3E2", "Dealers \u2192 Map \u2192 Drill"),
        ("\U0001F4C4", "Contract \u2192 Cancellation"),
        ("\U0001F916", "AI: \"find dealers in Quebec\""),
        ("\U0001F4A1", "Request a Feature"),
    ]
    card_w = Inches(3.95)
    card_h = Inches(2.0)
    gap_x = Inches(0.2)
    gap_y = Inches(0.3)
    total_w = card_w * 3 + gap_x * 2
    start_left = (SLIDE_W - total_w) / 2
    start_top = Inches(2.2)

    colors = [PRIMARY, ACCENT_TEAL, ACCENT_ORANGE,
              ACCENT_PINK, PRIMARY_DARK, ACCENT_GREEN]
    for i, (emoji, label) in enumerate(steps):
        row = i // 3
        col = i % 3
        left = start_left + (card_w + gap_x) * col
        top = start_top + (card_h + gap_y) * row
        add_rounded(slide, left, top, card_w, card_h, fill=WHITE,
                    line_color=BG_GRAY, line_width=Pt(0.75), corner=0.08)
        add_circle(slide, left + Inches(0.7), top + Inches(0.7),
                   Inches(0.45), fill=colors[i])
        add_text(slide, str(i + 1),
                 left + Inches(0.25), top + Inches(0.35),
                 Inches(0.9), Inches(0.7),
                 size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_emoji(slide, emoji,
                  left + card_w - Inches(0.95), top + Inches(0.25),
                  Inches(0.8), Inches(0.8), size=32)
        add_text(slide, label,
                 left + Inches(1.25), top + Inches(1.15),
                 card_w - Inches(1.4), Inches(0.7),
                 size=16, bold=True, color=TEXT_DARK)

    add_pill(slide, Inches(3.0), Inches(6.7), Inches(7.3), Inches(0.55),
             "Live app \u00b7 screenshots ready as a safety net",
             fill=PRIMARY_DARK, size=12)

    add_footer(slide, page, total)


def slide_thanks(prs, page=None, total=None):  # noqa: ARG001
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = PRIMARY_DARK

    add_circle(slide, Inches(1.5), Inches(1.5), Inches(0.5), fill=ACCENT_PINK)
    add_circle(slide, Inches(12.0), Inches(1.2), Inches(0.7), fill=ACCENT_TEAL)
    add_circle(slide, Inches(2.4), Inches(6.4), Inches(0.4), fill=ACCENT_YELLOW)
    add_circle(slide, Inches(11.5), Inches(6.5), Inches(0.55), fill=PRIMARY_LIGHT)

    add_text(slide, "Thank you \U0001F64C",
             Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6),
             size=72, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    add_text(slide, "Admin-only.  Secure.  Self-service.  AI inside.",
             Inches(0.9), Inches(4.1), Inches(11.5), Inches(0.7),
             size=24, color=PRIMARY_LIGHT, align=PP_ALIGN.CENTER)

    add_pill(slide, Inches(4.9), Inches(5.5), Inches(3.5), Inches(0.7),
             "Questions?  \u2192  Let's chat", fill=ACCENT_PINK, size=18)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    builders = [
        slide_title,           # 1
        slide_problem,         # 2
        slide_vision,          # 3
        slide_why_now,         # 4
        slide_what,            # 5
        slide_better_way,      # 6
        slide_before_after,    # 7
        slide_dashboard,       # 8
        slide_ecosystem,       # 9
        slide_ai,              # 10
        slide_feedback,        # 11
        slide_impact,          # 12
        slide_roadmap,         # 13
        slide_demo,            # 14
        slide_thanks,          # 15
    ]
    total = len(builders)

    for idx, build in enumerate(builders, start=1):
        if build in (slide_title, slide_thanks):
            build(prs)
        else:
            build(prs, idx, total)

    out_dir = Path(__file__).resolve().parent
    out_path = out_dir / "BackOffice-Admin-Portal-Demo.pptx"
    prs.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
